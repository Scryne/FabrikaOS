"""
FDYS — CRM Servisi
Faz 2 — Ortak & Firma Yönetimi İş Mantığı

Excel import/export, genel arama motoru, CRM istatistik ve
rapor servisleri.
"""

import logging
import os
from datetime import datetime, timezone
from sqlalchemy import or_, func

from app.extensions import db

logger = logging.getLogger(__name__)


def crm_arama(arama_terimi, tip=None, limit=50):
    """
    CRM genel arama motoru.
    Ortak adı, TC, telefon, firma ünvanı, vergi no üzerinde arama yapar.

    Args:
        arama_terimi: Arama metni
        tip: 'ortak' veya 'firma' veya None (tümü)
        limit: Maksimum sonuç sayısı

    Returns:
        dict: {'ortaklar': [...], 'firmalar': [...]}
    """
    from app.models.crm import Ortak, Firma

    sonuclar = {'ortaklar': [], 'firmalar': []}

    if not arama_terimi or len(arama_terimi) < 2:
        return sonuclar

    terim = f'%{arama_terimi}%'

    if tip is None or tip == 'ortak':
        ortaklar = Ortak.query.filter(
            Ortak.aktif == True,
            or_(
                Ortak.ad_soyad.ilike(terim),
                Ortak.tc_kimlik.ilike(terim),
                Ortak.telefon.ilike(terim),
                Ortak.telefon2.ilike(terim),
                Ortak.eposta.ilike(terim),
                Ortak.ortak_no.ilike(terim),
                Ortak.blok.ilike(terim),
                Ortak.daire_no.ilike(terim),
            )
        ).limit(limit).all()
        sonuclar['ortaklar'] = ortaklar

    if tip is None or tip == 'firma':
        firmalar = Firma.query.filter(
            Firma.aktif == True,
            or_(
                Firma.unvan.ilike(terim),
                Firma.vergi_no.ilike(terim),
                Firma.telefon.ilike(terim),
                Firma.eposta.ilike(terim),
                Firma.firma_no.ilike(terim),
                Firma.yetkili_kisi.ilike(terim),
                Firma.blok.ilike(terim),
                Firma.daire_no.ilike(terim),
            )
        ).limit(limit).all()
        sonuclar['firmalar'] = firmalar

    return sonuclar


def crm_dashboard_istatistik():
    """
    CRM dashboard için istatistik verisi toplar.

    Returns:
        dict: Dashboard KPI verileri
    """
    from app.models.crm import Ortak, Firma, FirmaYetkilisi, CrmIletisimGecmisi

    # Toplam ortak ve firma sayıları
    toplam_ortak = Ortak.query.filter_by(aktif=True).count()
    aktif_ortak = Ortak.query.filter_by(aktif=True, durum='aktif').count()
    toplam_firma = Firma.query.filter_by(aktif=True).count()
    aktif_firma = Firma.query.filter_by(aktif=True, durum='aktif').count()

    # Ortak tipi dağılımı
    ortak_tipi_dagilimi = db.session.query(
        Ortak.ortak_tipi, func.count(Ortak.id)
    ).filter(Ortak.aktif == True).group_by(Ortak.ortak_tipi).all()

    # Firma sektör dağılımı
    sektor_dagilimi = db.session.query(
        Firma.sektor, func.count(Firma.id)
    ).filter(Firma.aktif == True).group_by(Firma.sektor).all()

    # Son iletişim kayıtları
    son_iletisimler = CrmIletisimGecmisi.query.filter_by(aktif=True).order_by(
        CrmIletisimGecmisi.tarih.desc()
    ).limit(10).all()

    # Bu ay yapılan iletişim sayısı
    bugun = datetime.now(timezone.utc)
    bu_ay_iletisim = CrmIletisimGecmisi.query.filter(
        CrmIletisimGecmisi.aktif == True,
        func.extract('month', CrmIletisimGecmisi.tarih) == bugun.month,
        func.extract('year', CrmIletisimGecmisi.tarih) == bugun.year,
    ).count()

    # Blok bazlı ortak dağılımı
    blok_dagilimi = db.session.query(
        Ortak.blok, func.count(Ortak.id)
    ).filter(Ortak.aktif == True, Ortak.blok.isnot(None)).group_by(Ortak.blok).all()

    return {
        'toplam_ortak': toplam_ortak,
        'aktif_ortak': aktif_ortak,
        'toplam_firma': toplam_firma,
        'aktif_firma': aktif_firma,
        'ortak_tipi_dagilimi': [list(r) for r in ortak_tipi_dagilimi],
        'sektor_dagilimi': [list(r) for r in sektor_dagilimi],
        'son_iletisimler': son_iletisimler,
        'bu_ay_iletisim': bu_ay_iletisim,
        'blok_dagilimi': [list(r) for r in blok_dagilimi],
    }


def ortak_detay_rapor(ortak_id):
    """
    Ortak detay kartı için tüm verileri toplar.

    Args:
        ortak_id: Ortak ID

    Returns:
        dict: Ortak detay raporu veya None
    """
    from app.models.crm import Ortak, CrmIletisimGecmisi

    ortak = db.session.get(Ortak, ortak_id)
    if not ortak:
        return None

    # İletişim geçmişi
    iletisimler = CrmIletisimGecmisi.query.filter_by(
        ilgili_tip='ortak', ilgili_id=ortak_id, aktif=True
    ).order_by(CrmIletisimGecmisi.tarih.desc()).limit(20).all()

    return {
        'ortak': ortak,
        'iletisimler': iletisimler,
        'iletisim_sayisi': len(iletisimler),
    }


def firma_detay_rapor(firma_id):
    """
    Firma detay kartı için tüm verileri toplar.

    Args:
        firma_id: Firma ID

    Returns:
        dict: Firma detay raporu veya None
    """
    from app.models.crm import Firma, FirmaYetkilisi, CrmIletisimGecmisi

    firma = db.session.get(Firma, firma_id)
    if not firma:
        return None

    # Yetkililer
    yetkililer = FirmaYetkilisi.query.filter_by(
        firma_id=firma_id, aktif=True
    ).all()

    # İletişim geçmişi
    iletisimler = CrmIletisimGecmisi.query.filter_by(
        ilgili_tip='firma', ilgili_id=firma_id, aktif=True
    ).order_by(CrmIletisimGecmisi.tarih.desc()).limit(20).all()

    return {
        'firma': firma,
        'yetkililer': yetkililer,
        'iletisimler': iletisimler,
        'yetkili_sayisi': len(yetkililer),
        'iletisim_sayisi': len(iletisimler),
    }


def excel_export_ortaklar():
    """
    Tüm ortakları Excel formatında dışa aktarır.

    Returns:
        bytes: Excel dosya içeriği
    """
    from app.models.crm import Ortak
    import io

    try:
        import openpyxl
    except ImportError:
        logger.warning("openpyxl yüklü değil, Excel export yapılamıyor.")
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Ortaklar'

    # Başlık satırı
    basliklar = ['Ortak No', 'Ad Soyad', 'TC Kimlik', 'Telefon', 'Telefon 2',
                 'E-posta', 'Blok', 'Daire No', 'Ortak Tipi', 'Hisse (%)',
                 'Alan (m²)', 'Milliyet', 'Durum', 'Adres', 'Notlar']
    ws.append(basliklar)

    # Başlık stili
    from openpyxl.styles import Font, PatternFill
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='3B82F6', end_color='3B82F6', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    # Veri
    ortaklar = Ortak.query.filter_by(aktif=True).order_by(Ortak.ortak_no).all()
    for o in ortaklar:
        ws.append([
            o.ortak_no, o.ad_soyad, o.tc_kimlik or '', o.telefon or '',
            o.telefon2 or '', o.eposta or '', o.blok or '', o.daire_no or '',
            o.ortak_tipi_adi, float(o.hisse_orani) if o.hisse_orani else '',
            float(o.m2) if o.m2 else '', o.milliyet or '', o.durum_adi,
            o.adres or '', o.notlar or '',
        ])

    # Sütun genişlikleri otomatik ayarla
    for col in ws.columns:
        max_len = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_len:
                    max_len = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_len + 2, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def excel_export_firmalar():
    """
    Tüm firmaları Excel formatında dışa aktarır.

    Returns:
        bytes: Excel dosya içeriği
    """
    from app.models.crm import Firma
    import io

    try:
        import openpyxl
    except ImportError:
        logger.warning("openpyxl yüklü değil, Excel export yapılamıyor.")
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Firmalar'

    basliklar = ['Firma No', 'Ünvan', 'Vergi No', 'Vergi Dairesi', 'Sektör',
                 'Faaliyet Alanı', 'Yetkili Kişi', 'Telefon', 'E-posta',
                 'Web Sitesi', 'Blok', 'Daire No', 'Alan (m²)',
                 'Çalışan Sayısı', 'Durum', 'Adres', 'Notlar']
    ws.append(basliklar)

    from openpyxl.styles import Font, PatternFill
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='10B981', end_color='10B981', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    firmalar = Firma.query.filter_by(aktif=True).order_by(Firma.firma_no).all()
    for f in firmalar:
        ws.append([
            f.firma_no, f.unvan, f.vergi_no or '', f.vergi_dairesi or '',
            f.sektor_adi, f.faaliyet_alani or '', f.yetkili_kisi or '',
            f.telefon or '', f.eposta or '', f.web_sitesi or '',
            f.blok or '', f.daire_no or '',
            float(f.m2) if f.m2 else '', f.calisan_sayisi or '',
            f.durum_adi, f.adres or '', f.notlar or '',
        ])

    for col in ws.columns:
        max_len = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_len:
                    max_len = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_len + 2, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def excel_import_ortaklar(dosya_yolu):
    """
    Excel dosyasından ortakları toplu içe aktarır.

    Args:
        dosya_yolu: Yüklenen Excel dosya yolu

    Returns:
        dict: {'basarili': int, 'hatali': int, 'hatalar': list}
    """
    from app.models.crm import Ortak

    try:
        import openpyxl
    except ImportError:
        return {'basarili': 0, 'hatali': 0, 'hatalar': ['openpyxl yüklü değil.']}

    sonuc = {'basarili': 0, 'hatali': 0, 'hatalar': []}

    try:
        wb = openpyxl.load_workbook(dosya_yolu, read_only=True)
        ws = wb.active

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not row[0]:
                continue

            try:
                ortak_no = str(row[0]).strip()
                ad_soyad = str(row[1]).strip() if row[1] else None

                if not ad_soyad:
                    sonuc['hatalar'].append(f'Satır {row_idx}: Ad soyad boş.')
                    sonuc['hatali'] += 1
                    continue

                # Mevcut ortak kontrolü
                mevcut = Ortak.query.filter_by(ortak_no=ortak_no).first()
                if mevcut:
                    sonuc['hatalar'].append(f'Satır {row_idx}: Ortak no "{ortak_no}" zaten mevcut.')
                    sonuc['hatali'] += 1
                    continue

                ortak = Ortak(
                    ortak_no=ortak_no,
                    ad_soyad=ad_soyad,
                    tc_kimlik=str(row[2]).strip() if row[2] else None,
                    telefon=str(row[3]).strip() if row[3] else None,
                    telefon2=str(row[4]).strip() if len(row) > 4 and row[4] else None,
                    eposta=str(row[5]).strip() if len(row) > 5 and row[5] else None,
                    blok=str(row[6]).strip() if len(row) > 6 and row[6] else None,
                    daire_no=str(row[7]).strip() if len(row) > 7 and row[7] else None,
                    ortak_tipi=str(row[8]).strip().lower() if len(row) > 8 and row[8] else 'malik',
                    durum='aktif',
                )
                db.session.add(ortak)
                sonuc['basarili'] += 1

            except Exception as e:
                sonuc['hatalar'].append(f'Satır {row_idx}: {str(e)}')
                sonuc['hatali'] += 1

        db.session.commit()
        wb.close()

    except Exception as e:
        db.session.rollback()
        sonuc['hatalar'].append(f'Dosya okuma hatası: {str(e)}')
        logger.error(f"Excel import hatası: {e}")

    return sonuc


def excel_import_firmalar(dosya_yolu):
    """
    Excel dosyasından firmaları toplu içe aktarır.

    Args:
        dosya_yolu: Yüklenen Excel dosya yolu

    Returns:
        dict: {'basarili': int, 'hatali': int, 'hatalar': list}
    """
    from app.models.crm import Firma

    try:
        import openpyxl
    except ImportError:
        return {'basarili': 0, 'hatali': 0, 'hatalar': ['openpyxl yüklü değil.']}

    sonuc = {'basarili': 0, 'hatali': 0, 'hatalar': []}

    try:
        wb = openpyxl.load_workbook(dosya_yolu, read_only=True)
        ws = wb.active

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not row[0]:
                continue

            try:
                firma_no = str(row[0]).strip()
                unvan = str(row[1]).strip() if row[1] else None

                if not unvan:
                    sonuc['hatalar'].append(f'Satır {row_idx}: Ünvan boş.')
                    sonuc['hatali'] += 1
                    continue

                mevcut = Firma.query.filter_by(firma_no=firma_no).first()
                if mevcut:
                    sonuc['hatalar'].append(f'Satır {row_idx}: Firma no "{firma_no}" zaten mevcut.')
                    sonuc['hatali'] += 1
                    continue
                    
                vergi_no = str(row[2]).strip() if len(row) > 2 and row[2] else None
                if vergi_no:
                    mevcut_vn = Firma.query.filter_by(vergi_no=vergi_no).first()
                    if mevcut_vn:
                        sonuc['hatalar'].append(f'Satır {row_idx}: Vergi no "{vergi_no}" zaten mevcut.')
                        sonuc['hatali'] += 1
                        continue

                firma = Firma(
                    firma_no=firma_no,
                    unvan=unvan,
                    vergi_no=vergi_no,
                    vergi_dairesi=str(row[3]).strip() if len(row) > 3 and row[3] else None,
                    sektor=str(row[4]).strip().lower() if len(row) > 4 and row[4] else 'diger',
                    faaliyet_alani=str(row[5]).strip() if len(row) > 5 and row[5] else None,
                    yetkili_kisi=str(row[6]).strip() if len(row) > 6 and row[6] else None,
                    telefon=str(row[7]).strip() if len(row) > 7 and row[7] else None,
                    eposta=str(row[8]).strip() if len(row) > 8 and row[8] else None,
                    blok=str(row[9]).strip() if len(row) > 9 and row[9] else None,
                    daire_no=str(row[10]).strip() if len(row) > 10 and row[10] else None,
                    durum='aktif',
                )
                db.session.add(firma)
                sonuc['basarili'] += 1

            except Exception as e:
                sonuc['hatalar'].append(f'Satır {row_idx}: {str(e)}')
                sonuc['hatali'] += 1

        db.session.commit()
        wb.close()

    except Exception as e:
        db.session.rollback()
        sonuc['hatalar'].append(f'Dosya okuma hatası: {str(e)}')
        logger.error(f"Firma excel import hatası: {e}")

    return sonuc


def harita_verileri():
    """
    CRM harita görünümü için ortak ve firma GPS verilerini toplar.

    Returns:
        dict: {'ortaklar': [...], 'firmalar': [...]}
    """
    from app.models.crm import Ortak, Firma

    ortaklar = Ortak.query.filter(
        Ortak.aktif == True,
        Ortak.harita_enlem.isnot(None),
        Ortak.harita_boylam.isnot(None),
    ).all()

    firmalar = Firma.query.filter(
        Firma.aktif == True,
        Firma.harita_enlem.isnot(None),
        Firma.harita_boylam.isnot(None),
    ).all()

    return {
        'ortaklar': [{
            'id': o.id,
            'ad_soyad': o.ad_soyad,
            'ortak_no': o.ortak_no,
            'blok': o.blok or '',
            'daire_no': o.daire_no or '',
            'ortak_tipi': o.ortak_tipi_adi,
            'telefon': o.telefon or '',
            'durum': o.durum,
            'lat': o.harita_enlem,
            'lng': o.harita_boylam,
        } for o in ortaklar],
        'firmalar': [{
            'id': f.id,
            'unvan': f.unvan,
            'firma_no': f.firma_no,
            'blok': f.blok or '',
            'daire_no': f.daire_no or '',
            'sektor': f.sektor_adi,
            'telefon': f.telefon or '',
            'durum': f.durum,
            'lat': f.harita_enlem,
            'lng': f.harita_boylam,
        } for f in firmalar],
    }


def sonraki_ortak_no():
    """
    Sıradaki ortak numarasını üretir.

    Returns:
        str: Yeni ortak no (ör: "ORT-0001")
    """
    from app.models.crm import Ortak
    son = Ortak.query.order_by(Ortak.id.desc()).first()
    if son:
        try:
            # ORT-0001 formatı
            sayi = int(son.ortak_no.split('-')[-1]) + 1
        except (ValueError, IndexError):
            sayi = son.id + 1
    else:
        sayi = 1
    return f'ORT-{sayi:04d}'


def sonraki_firma_no():
    """
    Sıradaki firma numarasını üretir.

    Returns:
        str: Yeni firma no (ör: "FRM-0001")
    """
    from app.models.crm import Firma
    son = Firma.query.order_by(Firma.id.desc()).first()
    if son:
        try:
            sayi = int(son.firma_no.split('-')[-1]) + 1
        except (ValueError, IndexError):
            sayi = son.id + 1
    else:
        sayi = 1
    return f'FRM-{sayi:04d}'
