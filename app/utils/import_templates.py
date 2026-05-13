"""
FDYS — Excel Şablon Oluşturucu
Faz 5 — Import için boş şablon dosyaları
"""

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation


def _stil_uygula(ws, basliklar):
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    for col_idx, baslik in enumerate(basliklar, 1):
        cell = ws.cell(row=1, column=col_idx, value=baslik)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = max(len(baslik) + 4, 15)


def sablon_olustur(tip: str) -> BytesIO:
    """Belirtilen veri tipi için import şablonu oluşturur."""
    wb = Workbook()
    ws = wb.active

    if tip == 'malzeme':
        ws.title = "Malzeme Import"
        basliklar = ['Kod', 'Ad', 'Birim', 'Kategori', 'Lokasyon', 'Birim Fiyat', 'Minimum Stok', 'Mevcut Miktar']
        _stil_uygula(ws, basliklar)
        # Örnek veri
        ws.append(['MLZ-001', 'Rulman 6205', 'adet', 'Yedek Parça', 'A-1 Rafı', 45.50, 10, 150])
        ws.append(['MLZ-002', 'Hidrolik Yağ 46', 'litre', 'Sarf Malzeme', 'B-3 Depo', 120.00, 50, 200])
        # Birim dropdown
        dv = DataValidation(type="list", formula1='"adet,kg,litre,metre"', allow_blank=True)
        dv.prompt = "Birim seçin"
        ws.add_data_validation(dv)
        dv.add(f'C2:C1000')

    elif tip == 'calisan':
        ws.title = "Çalışan Import"
        basliklar = ['Sicil No', 'Ad Soyad', 'Departman', 'Pozisyon', 'İşe Başlama', 'Telefon', 'E-posta', 'Ücret']
        _stil_uygula(ws, basliklar)
        ws.append(['P-001', 'Ahmet Yılmaz', 'Üretim', 'Operatör', '15.01.2024', '0532-111-2233', 'ahmet@firma.com', 15000])
        ws.append(['P-002', 'Ayşe Demir', 'Kalite', 'Kalite Uzmanı', '01.03.2024', '0533-222-3344', 'ayse@firma.com', 18000])

    elif tip == 'makine':
        ws.title = "Makine Import"
        basliklar = ['Kod', 'Ad', 'Marka', 'Model', 'Teorik Kapasite', 'Satın Alma Tarihi', 'Lokasyon', 'Hedef OEE', 'Bakım Periyodu (Gün)']
        _stil_uygula(ws, basliklar)
        ws.append(['CNC-001', 'CNC Torna T-200', 'Haas', 'ST-20', 100, '01.06.2022', 'A Hattı', 85, 90])
        ws.append(['PRE-001', 'Hidrolik Pres 250T', 'Ermaksan', 'HP-250', 80, '15.09.2023', 'B Hattı', 80, 60])
    else:
        raise ValueError(f"Geçersiz şablon tipi: {tip}")

    # Örnek veri satırlarını farklı renkle işaretle
    ornek_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    ornek_font = Font(italic=True, color="856404", name="Calibri", size=10)
    for row in ws.iter_rows(min_row=2, max_row=3, max_col=ws.max_column):
        for cell in row:
            cell.fill = ornek_fill
            cell.font = ornek_font

    # Açıklama notu ekle
    not_row = ws.max_row + 2
    ws.cell(row=not_row, column=1, value="ℹ️ Sarı satırlar örnek verilerdir. Import öncesinde siliniz veya üzerine yazınız.")
    ws.cell(row=not_row, column=1).font = Font(italic=True, color="6C757D", size=9)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
