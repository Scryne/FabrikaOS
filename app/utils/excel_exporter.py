"""FDYS — Excel Rapor Dışa Aktarma Servisi (Kurumsal Sürüm)"""

from io import BytesIO
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

class ExcelExporter:
    def __init__(self):
        self.header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        self.header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
        self.center_alignment = Alignment(horizontal="center", vertical="center")
        self.left_alignment = Alignment(horizontal="left", vertical="center")

    def _apply_header_style(self, ws, row_idx=1):
        for cell in ws[row_idx]:
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_alignment

    def _auto_adjust_columns(self, ws):
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass  # Hücre değeri ölçüm hatası, varsayılan genişlik kullanılır
            adjusted_width = (max_length + 2)
            if adjusted_width > 50:
                adjusted_width = 50 # Cap width
            ws.column_dimensions[column].width = adjusted_width

    def listeyi_disa_aktar(self, basliklar, veriler, sayfa_adi="Rapor"):
        """Genel amaçlı tek sayfalık dışa aktarıcı."""
        wb = Workbook()
        ws = wb.active
        ws.title = sayfa_adi

        # Başlıkları Yaz
        ws.append(basliklar)
        self._apply_header_style(ws)

        # Verileri Yaz
        for satir in veriler:
            ws.append(satir)
            
        # Hücre formatlama (Genel hizalama)
        for row in ws.iter_rows(min_row=2, max_col=len(basliklar), max_row=ws.max_row):
            for cell in row:
                cell.alignment = self.left_alignment

        self._auto_adjust_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def cok_sayfali_rapor_aktar(self, sayfalar_dict):
        """
        sayfalar_dict formatı:
        {
            "Sayfa 1": {"basliklar": ["A", "B"], "veriler": [[1, 2], [3, 4]]},
            "Sayfa 2": {"basliklar": ["C", "D"], "veriler": [[5, 6], [7, 8]]}
        }
        """
        wb = Workbook()
        
        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])
            
        for sayfa_adi, icerik in sayfalar_dict.items():
            ws = wb.create_sheet(title=sayfa_adi)
            ws.append(icerik["basliklar"])
            self._apply_header_style(ws)
            
            for satir in icerik["veriler"]:
                ws.append(satir)
                
            self._auto_adjust_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
